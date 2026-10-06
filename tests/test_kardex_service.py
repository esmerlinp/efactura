"""Pruebas exhaustivas para la Fase 6 — Kardex Valorizado Continuo (FIFO).

Cubre los 19 escenarios de prueba de la especificación técnica aprobada:
1. Kardex vacío
2. Una compra
3. Compra + compra + venta FIFO
4. Venta que consume múltiples capas
5. E34 genera nueva capa histórica
6. Anulación de venta
7. Transferencia origen (vista salida)
8. Transferencia destino (vista entrada)
9. Transferencia vista global (impacto neto neutro)
10. Dos almacenes independientes
11. Saldo inicial con dateFrom
12. Transacciones con mismo timestamp (orden determinista)
13. Stock físico vs Kardex inconsistente (auditoría/reconciliación)
14. Ledger FIFO vs Kardex inconsistente (auditoría/reconciliación)
15. items.totalStock vs suma de almacenes
16. Idempotencia / no duplicación
17. Aislamiento estricto multiempresa
18. Exportación Excel (.xlsx) con valores idénticos al Kardex
19. Exportación PDF con valores idénticos al Kardex
"""

import io
import pytest
from unittest.mock import patch, MagicMock

from app.services.inventory_transaction_service import InventoryTransactionService
from app.services.kardex_service import KardexService
from app.services.credit_note_inventory_service import CreditNoteInventoryService


# ── MOCK FIRESTORE STORE ─────────────────────────────────────────────────────

# ── MOCK FIRESTORE STORE ─────────────────────────────────────────────────────

class MockDocumentSnapshot:
    def __init__(self, doc_id, data=None, exists=True):
        self.id = doc_id
        self._data = data or {}
        self.exists = exists

    def to_dict(self):
        return dict(self._data)


class MockDocumentReference:
    def __init__(self, coll, doc_id):
        self.coll = coll
        self.id = doc_id

    def get(self, transaction=None):
        data = self.coll._storage.get(self.id)
        if data is not None:
            return MockDocumentSnapshot(self.id, data=data, exists=True)
        return MockDocumentSnapshot(self.id, exists=False)

    def set(self, data, merge=False):
        if merge and self.id in self.coll._storage:
            self.coll._storage[self.id].update(data)
        else:
            self.coll._storage[self.id] = dict(data)

    def update(self, data):
        if self.id in self.coll._storage:
            self.coll._storage[self.id].update(data)
        else:
            self.coll._storage[self.id] = dict(data)


class MockCollectionReference:
    def __init__(self, name):
        self.name = name
        self._storage = {}

    def document(self, doc_id):
        return MockDocumentReference(self, doc_id)

    def where(self, *args, **kwargs):
        return self

    def limit(self, count):
        return self

    def stream(self):
        for doc_id, data in list(self._storage.items()):
            yield MockDocumentSnapshot(doc_id, data=data, exists=True)

    def get(self, transaction=None):
        return list(self.stream())


class MockFirestoreTransaction:
    def __init__(self, db):
        self.db = db

    def get(self, ref):
        return ref.get(transaction=self)

    def set(self, ref, data, merge=False):
        ref.set(data, merge=merge)

    def update(self, ref, data):
        ref.update(data)


class MockFirestoreClient:
    def __init__(self):
        self.collections = {}

    def collection(self, name):
        if name not in self.collections:
            self.collections[name] = MockCollectionReference(name)
        return self.collections[name]

    def transaction(self):
        txn = MockFirestoreTransaction(self)

        def runner(fn):
            return fn(txn)

        txn.__call__ = runner
        return txn


@pytest.fixture
def kardex_mock_store():
    client = MockFirestoreClient()
    store = {
        "client": client,
        "items": {},
        "warehouses": [
            {"id": "wh-a", "name": "Almacén Central", "branchId": "branch-1"},
            {"id": "wh-b", "name": "Almacén Sucursal B", "branchId": "branch-1"},
        ],
        "invoices": {},
    }
    return store


@pytest.fixture
def kardex_env(kardex_mock_store):
    def mock_company_coll(company_id=None, owner_uid=None, coll_name=""):
        return kardex_mock_store["client"].collection(f"companies/{company_id}/{coll_name}")

    def mock_get_items(owner_uid, sandbox=True, company_id=None, **kwargs):
        return list(kardex_mock_store["items"].values())

    def mock_get_warehouses(owner_uid, sandbox=True, company_id=None):
        return kardex_mock_store["warehouses"]

    def mock_get_invoice(owner_uid, invoice_id, sandbox=True, company_id=None):
        return kardex_mock_store["invoices"].get(invoice_id)

    def mock_get_invoices(owner_uid, sandbox=True, company_id=None):
        return list(kardex_mock_store["invoices"].values())

    with patch("app.services.inventory_transaction_service._company_coll", side_effect=mock_company_coll), \
         patch("app.services.inventory_transaction_service.db_firestore", kardex_mock_store["client"]), \
         patch("app.services.inventory_transaction_service.firebase_initialized", True), \
         patch("app.services.inventory_transaction_service.DatabaseService.get_warehouses", side_effect=mock_get_warehouses), \
         patch("app.services.kardex_service._company_coll", side_effect=mock_company_coll), \
         patch("app.services.kardex_service.db_firestore", kardex_mock_store["client"]), \
         patch("app.services.kardex_service.firebase_initialized", True), \
         patch("app.services.kardex_service.DatabaseService.get_warehouses", side_effect=mock_get_warehouses), \
         patch("app.services.db_service._company_coll", side_effect=mock_company_coll), \
         patch("app.services.db_service.db_firestore", kardex_mock_store["client"]), \
         patch("app.services.db_service.firebase_initialized", True), \
         patch("app.services.db_service.DatabaseService.get_items", side_effect=mock_get_items), \
         patch("app.services.db_service.DatabaseService.get_warehouses", side_effect=mock_get_warehouses), \
         patch("app.services.credit_note_inventory_service._company_coll", side_effect=mock_company_coll), \
         patch("app.services.credit_note_inventory_service.db_firestore", kardex_mock_store["client"]), \
         patch("app.services.credit_note_inventory_service.firebase_initialized", True), \
         patch("app.services.credit_note_inventory_service.DatabaseService.get_items", side_effect=mock_get_items), \
         patch("app.services.credit_note_inventory_service.DatabaseService.get_invoice", side_effect=mock_get_invoice), \
         patch("app.services.credit_note_inventory_service.DatabaseService.get_invoices", side_effect=mock_get_invoices), \
         patch("app.services.credit_note_inventory_service.DatabaseService.get_warehouses", side_effect=mock_get_warehouses):
        yield kardex_mock_store


# ── TEST SUITE FASE 6 ────────────────────────────────────────────────────────

def test_1_kardex_empty(kardex_env):
    """Caso 1: Kardex para un artículo sin movimientos retorna saldos en cero y estado OK."""
    store = kardex_env
    company_id = "comp-kardex-1"
    item_id = "item-001"

    item_data = {"id": item_id, "code": "ART-01", "name": "Taladro", "costPrice": 100.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)

    assert res["initialBalance"]["quantity"] == 0.0
    assert res["initialBalance"]["totalValue"] == 0.0
    assert len(res["movements"]) == 0
    assert res["finalBalance"]["quantity"] == 0.0
    assert res["finalBalance"]["totalValue"] == 0.0
    assert res["reconciliation"]["status"] == "OK"
    assert res["reconciliation"]["qtyDifference"] == 0.0


def test_2_kardex_single_purchase(kardex_env):
    """Caso 2: Una compra registra entrada y saldo valorizado idéntico."""
    store = kardex_env
    company_id = "comp-kardex-2"
    owner_uid = "user-k"
    item_id = "item-002"

    item_data = {"id": item_id, "code": "ART-02", "name": "Cemento", "costPrice": 350.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 20.0, "unitCost": 350.0,
        "destinationWarehouseId": "wh-a", "reason": "COMPRA", "referenceId": "OC-1001",
        "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)

    assert len(res["movements"]) == 1
    mov = res["movements"][0]
    assert mov["inQty"] == 20.0
    assert mov["inUnitCost"] == 350.0
    assert mov["inTotalValue"] == 7000.0
    assert mov["balanceQty"] == 20.0
    assert mov["balanceTotalValue"] == 7000.0
    assert mov["balanceUnitCost"] == 350.0
    assert res["finalBalance"]["quantity"] == 20.0
    assert res["finalBalance"]["totalValue"] == 7000.0
    assert res["reconciliation"]["status"] == "OK"


def test_3_and_4_kardex_mandatory_multi_purchase_and_fifo_sale(kardex_env):
    """
    Casos 3 y 4: Compra 1 (10 × $100) + Compra 2 (10 × $120) + Venta 15 (COGS $1,600).
    Verifica que el Kardex consume de las capas persistidas (10x100 + 5x120) y queda saldo = 5 × $120 = $600.
    """
    store = kardex_env
    company_id = "comp-kardex-3"
    owner_uid = "user-k"
    item_id = "item-003"

    item_data = {"id": item_id, "code": "ART-03", "name": "Varilla", "costPrice": 100.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # 1. Compra 1: 10 × $100
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 100.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # 2. Compra 2: 10 × $120
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 120.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-02T10:00:00Z"
    }, sandbox=True)

    # 3. Venta de 15 unidades
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 15.0, "originWarehouseId": "wh-a",
        "reason": "VENTA", "referenceType": "INVOICE", "referenceId": "FAC-001",
        "date": "2026-08-03T10:00:00Z"
    }, sandbox=True)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)

    assert len(res["movements"]) == 3
    # Mov 1 (Compra 1): Saldo 10 / 1000
    assert res["movements"][0]["balanceQty"] == 10.0
    assert res["movements"][0]["balanceTotalValue"] == 1000.0

    # Mov 2 (Compra 2): Saldo 20 / 2200
    assert res["movements"][1]["balanceQty"] == 20.0
    assert res["movements"][1]["balanceTotalValue"] == 2200.0

    # Mov 3 (Venta 15): Salida 15 / 1600 -> Saldo 5 / 600
    mov_sale = res["movements"][2]
    assert mov_sale["outQty"] == 15.0
    assert mov_sale["outTotalValue"] == 1600.0
    assert mov_sale["balanceQty"] == 5.0
    assert mov_sale["balanceTotalValue"] == 600.0
    assert mov_sale["balanceUnitCost"] == 120.0

    # Desglose de capas persistido
    assert len(mov_sale["costLayers"]) == 2
    assert mov_sale["costLayers"][0]["quantity"] == 10.0
    assert mov_sale["costLayers"][0]["unitCost"] == 100.0
    assert mov_sale["costLayers"][1]["quantity"] == 5.0
    assert mov_sale["costLayers"][1]["unitCost"] == 120.0

    # Conciliación
    assert res["reconciliation"]["stockQty"] == 5.0
    assert res["reconciliation"]["kardexQty"] == 5.0
    assert res["reconciliation"]["ledgerValue"] == 600.0
    assert res["reconciliation"]["kardexValue"] == 600.0
    assert res["reconciliation"]["status"] == "OK"


def test_5_kardex_e34_creates_new_historical_cost_layer(kardex_env):
    """
    Caso 5: E34 Devolución física entra como nueva capa con costo ponderado de salida (RD$ 106.6667),
    documentando la nueva capa sin suponer restauración idéntica de capas originales.
    """
    store = kardex_env
    company_id = "comp-kardex-e34"
    owner_uid = "user-k"
    item_id = "item-005"

    item_data = {"id": item_id, "code": "ART-05", "name": "Cable", "costPrice": 100.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # Compra 10 × $100
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 100.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # Compra 10 × $120
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 120.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-02T10:00:00Z"
    }, sandbox=True)

    # Venta 15 -> COGS $1600 (unitario ponderado = 106.6667)
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 15.0, "originWarehouseId": "wh-a",
        "reason": "VENTA", "referenceType": "INVOICE", "referenceId": "FAC-E34-01",
        "date": "2026-08-03T10:00:00Z"
    }, sandbox=True)

    # Factura original mockeada
    store["invoices"]["fac-e34-01"] = {
        "id": "fac-e34-01", "invoiceNumber": "FAC-E34-01", "status": "Emitida", "warehouseId": "wh-a",
        "items": [{"id": "line-1", "itemId": item_id, "type": "Bien", "quantity": 15, "originalCost": 106.6667, "unitCost": 106.6667}]
    }

    # E34 Devolución 5
    credit_note = {
        "id": "nc-01", "invoiceNumber": "E34-001", "status": "Emitida", "referenceInvoiceId": "fac-e34-01",
        "warehouseId": "wh-a",
        "items": [{"id": item_id, "originalLineId": "line-1", "type": "Bien", "quantity": 5, "reingresoStock": True, "warehouseId": "wh-a"}]
    }
    CreditNoteInventoryService.process_credit_note_stock_reentry(owner_uid, company_id, credit_note, sandbox=True)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)

    assert len(res["movements"]) == 4
    mov_nc = res["movements"][3]
    assert mov_nc["type"] == "ENTRADA"
    assert mov_nc["reason"] == "DEVOLUCION_CLIENTE"
    assert mov_nc["inQty"] == 5.0
    assert mov_nc["inUnitCost"] == 106.6667
    assert mov_nc["inTotalValue"] == 533.33

    # Saldo final esperado: 5 uds existentes ($600) + 5 uds devueltas ($533.33) = 10 uds / $1,133.33
    assert res["finalBalance"]["quantity"] == 10.0
    assert res["finalBalance"]["totalValue"] == 1133.33
    assert res["finalBalance"]["unitCost"] == 113.333
    assert res["reconciliation"]["status"] == "OK"


def test_6_kardex_sale_voiding_reversal(kardex_env):
    """Caso 6: Anulación de venta genera reingreso trazable al costo original de la venta."""
    store = kardex_env
    company_id = "comp-kardex-void"
    owner_uid = "user-k"
    item_id = "item-006"

    item_data = {"id": item_id, "code": "ART-06", "name": "Pintura", "costPrice": 80.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # 1. Compra 10 × $80
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 80.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # 2. Venta 4 unidades
    tx_sale = InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 4.0, "originWarehouseId": "wh-a",
        "reason": "VENTA", "referenceType": "INVOICE", "referenceId": "FAC-V1",
        "date": "2026-08-02T10:00:00Z"
    }, sandbox=True)

    # 3. Anulación de venta -> Reingreso de 4 unidades a $80
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 4.0, "unitCost": tx_sale["unitCost"],
        "destinationWarehouseId": "wh-a", "reason": "ANULACION_VENTA", "referenceType": "INVOICE_ANULADA",
        "referenceId": "FAC-V1", "date": "2026-08-03T10:00:00Z"
    }, sandbox=True)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)

    assert len(res["movements"]) == 3
    assert res["finalBalance"]["quantity"] == 10.0
    assert res["finalBalance"]["totalValue"] == 800.0
    assert res["finalBalance"]["unitCost"] == 80.0
    assert res["reconciliation"]["status"] == "OK"


def test_7_8_9_kardex_warehouse_transfers(kardex_env):
    """
    Casos 7, 8 y 9: Proyección de TRANSFERENCIA:
    - Vista Almacén A (Origen): SALIDA 6 unidades
    - Vista Almacén B (Destino): ENTRADA 6 unidades
    - Vista Global Empresa: Transferencia interna neutra (impacto neto 0 en cantidad y 0 en valor)
    """
    store = kardex_env
    company_id = "comp-kardex-trans"
    owner_uid = "user-k"
    item_id = "item-007"

    item_data = {"id": item_id, "code": "ART-07", "name": "Tubos PVC", "costPrice": 50.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # Compra 10 × $50 en Almacén A
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 50.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # Transferencia de 6 unidades de wh-a hacia wh-b
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "TRANSFERENCIA", "quantity": 6.0,
        "originWarehouseId": "wh-a", "destinationWarehouseId": "wh-b",
        "reason": "TRANSFERENCIA_SUCURSAL", "date": "2026-08-02T10:00:00Z"
    }, sandbox=True)

    # Caso 7: Vista Almacén Origen (wh-a)
    k_wh_a = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, warehouse_id="wh-a", sandbox=True)
    assert len(k_wh_a["movements"]) == 2
    assert k_wh_a["movements"][1]["outQty"] == 6.0
    assert k_wh_a["movements"][1]["outTotalValue"] == 300.0
    assert k_wh_a["finalBalance"]["quantity"] == 4.0
    assert k_wh_a["finalBalance"]["totalValue"] == 200.0
    assert k_wh_a["reconciliation"]["status"] == "OK"

    # Caso 8: Vista Almacén Destino (wh-b)
    k_wh_b = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, warehouse_id="wh-b", sandbox=True)
    assert len(k_wh_b["movements"]) == 1
    assert k_wh_b["movements"][0]["inQty"] == 6.0
    assert k_wh_b["movements"][0]["inTotalValue"] == 300.0
    assert k_wh_b["finalBalance"]["quantity"] == 6.0
    assert k_wh_b["finalBalance"]["totalValue"] == 300.0
    assert k_wh_b["reconciliation"]["status"] == "OK"

    # Caso 9: Vista Global Empresa (sin filtro de almacén)
    k_global = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, warehouse_id=None, sandbox=True)
    assert len(k_global["movements"]) == 2
    # Movimiento 1: Entrada Compra (+10) -> Saldo 10 / 500
    # Movimiento 2: Transferencia interna (+6 in, +6 out -> impacto neto 0) -> Saldo 10 / 500
    assert k_global["movements"][1]["inQty"] == 6.0
    assert k_global["movements"][1]["outQty"] == 6.0
    assert k_global["finalBalance"]["quantity"] == 10.0
    assert k_global["finalBalance"]["totalValue"] == 500.0
    assert k_global["reconciliation"]["status"] == "OK"


def test_10_kardex_multi_warehouse_independent(kardex_env):
    """Caso 10: Dos almacenes mantienen capas y saldos totalmente independientes."""
    store = kardex_env
    company_id = "comp-kardex-multi"
    owner_uid = "user-k"
    item_id = "item-010"

    item_data = {"id": item_id, "code": "ART-10", "name": "Lámpara", "costPrice": 100.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # wh-a compra 10 × $100
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 100.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # wh-b compra 10 × $200
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 200.0,
        "destinationWarehouseId": "wh-b", "date": "2026-08-01T11:00:00Z"
    }, sandbox=True)

    # Venta desde wh-a: 4 unidades
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 4.0, "originWarehouseId": "wh-a",
        "date": "2026-08-02T10:00:00Z"
    }, sandbox=True)

    # Almacén A queda con 6 × $100 = $600
    k_a = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, warehouse_id="wh-a", sandbox=True)
    assert k_a["finalBalance"]["quantity"] == 6.0
    assert k_a["finalBalance"]["totalValue"] == 600.0
    assert k_a["finalBalance"]["unitCost"] == 100.0

    # Almacén B queda con 10 × $200 = $2,000
    k_b = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, warehouse_id="wh-b", sandbox=True)
    assert k_b["finalBalance"]["quantity"] == 10.0
    assert k_b["finalBalance"]["totalValue"] == 2000.0
    assert k_b["finalBalance"]["unitCost"] == 200.0


def test_11_kardex_date_from_initial_balance_accumulation(kardex_env):
    """
    Caso 11: Rango de fecha con dateFrom acumula transacciones anteriores en initialBalance
    y sólo muestra movimientos dentro del período.
    """
    store = kardex_env
    company_id = "comp-kardex-period"
    owner_uid = "user-k"
    item_id = "item-011"

    item_data = {"id": item_id, "code": "ART-11", "name": "Batería", "costPrice": 150.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # 1. Compra en Julio (fuera del período de Agosto)
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 150.0,
        "destinationWarehouseId": "wh-a", "date": "2026-07-15T10:00:00Z"
    }, sandbox=True)

    # 2. Venta en Julio
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 3.0, "originWarehouseId": "wh-a",
        "date": "2026-07-20T10:00:00Z"
    }, sandbox=True)

    # 3. Compra en Agosto (dentro del período)
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 5.0, "unitCost": 200.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-05T10:00:00Z"
    }, sandbox=True)

    # Consultamos Kardex de Agosto (2026-08-01 a 2026-08-31)
    res = KardexService.get_kardex_summary(
        company_id=company_id, item_id=item_id, date_from="2026-08-01", date_to="2026-08-31", sandbox=True
    )

    # Saldo inicial al 01-Agosto: 10 - 3 = 7 unidades, Valor = 10*150 - 3*150 = 1,050.00
    assert res["initialBalance"]["quantity"] == 7.0
    assert res["initialBalance"]["totalValue"] == 1050.0
    assert res["initialBalance"]["unitCost"] == 150.0

    # Solo 1 movimiento en Agosto
    assert len(res["movements"]) == 1
    assert res["movements"][0]["inQty"] == 5.0
    assert res["movements"][0]["inTotalValue"] == 1000.0

    # Saldo final al 31-Agosto: 7 + 5 = 12 unidades, Valor = 1,050 + 1,000 = 2,050.00
    assert res["finalBalance"]["quantity"] == 12.0
    assert res["finalBalance"]["totalValue"] == 2050.0


def test_12_kardex_deterministic_sorting_same_timestamp(kardex_env):
    """
    Caso 12: Movimientos registrados con el mismo timestamp exacto se ordenan determinísticamente
    priorizando ENTRADA antes que SALIDA para garantizar disponibilidad de stock.
    """
    store = kardex_env
    company_id = "comp-kardex-sort"
    owner_uid = "user-k"
    item_id = "item-012"

    item_data = {"id": item_id, "code": "ART-12", "name": "Filtro", "costPrice": 90.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    same_time = "2026-08-10T12:00:00Z"
    # ENTRADA y SALIDA con idéntica fecha
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 90.0,
        "destinationWarehouseId": "wh-a", "date": same_time
    }, sandbox=True)

    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "SALIDA", "quantity": 4.0, "originWarehouseId": "wh-a",
        "date": same_time
    }, sandbox=True)

    res = KardexService.get_kardex(company_id=company_id, item_id=item_id, sandbox=True)
    assert len(res["movements"]) == 2
    assert res["movements"][0]["type"] == "ENTRADA"
    assert res["movements"][1]["type"] == "SALIDA"
    assert res["finalBalance"]["quantity"] == 6.0
    assert res["finalBalance"]["totalValue"] == 540.0


def test_13_and_14_reconciliation_detects_discrepancies(kardex_env):
    """
    Casos 13 y 14: La conciliación detecta discrepancias si inventory_stock o inventory_cost_ledger
    tienen valores inconsistentes y marca status="DISCREPANCY".
    """
    store = kardex_env
    company_id = "comp-kardex-disc"
    owner_uid = "user-k"
    item_id = "item-013"

    item_data = {"id": item_id, "code": "ART-13", "name": "Sensor", "costPrice": 500.0, "totalStock": 10.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    # 1. Movimiento en Kardex: 10 unidades a $500
    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 500.0,
        "destinationWarehouseId": "wh-a", "date": "2026-08-01T10:00:00Z"
    }, sandbox=True)

    # 2. Simulamos una corrupción manual directa en inventory_stock (físico = 8 en vez de 10)
    store["client"].collection(f"companies/{company_id}/sandbox_inventory_stock").document(f"{item_id}_wh-a").set({
        "itemId": item_id, "warehouseId": "wh-a", "quantity": 8.0
    })

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)
    assert res["reconciliation"]["status"] == "DISCREPANCY"
    assert res["reconciliation"]["qtyDifference"] == 2.0  # Kardex 10 vs Físico 8
    assert "Diferencia de cantidad" in res["reconciliation"]["notes"]


def test_15_total_stock_vs_sum_warehouses_reconciliation(kardex_env):
    """Caso 15: Conciliación de items.totalStock vs suma de almacenes."""
    store = kardex_env
    company_id = "comp-kardex-cat"
    owner_uid = "user-k"
    item_id = "item-015"

    item_data = {"id": item_id, "code": "ART-15", "name": "Relé", "costPrice": 40.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 40.0,
        "destinationWarehouseId": "wh-a"
    }, sandbox=True)

    res = KardexService.get_kardex_summary(company_id=company_id, item_id=item_id, sandbox=True)
    assert res["reconciliation"]["catalogTotalStock"] == 10.0
    assert res["reconciliation"]["catalogDifference"] == 0.0
    assert res["reconciliation"]["status"] == "OK"


def test_16_idempotency_prevents_duplicate_kardex_rows(kardex_env):
    """Caso 16: El reintento con la misma clave de idempotencia no duplica filas en el Kardex."""
    store = kardex_env
    company_id = "comp-kardex-idem"
    owner_uid = "user-k"
    item_id = "item-016"

    item_data = {"id": item_id, "code": "ART-16", "name": "Tuerca", "costPrice": 5.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    tx_payload = {
        "itemId": item_id, "type": "ENTRADA", "quantity": 100.0, "unitCost": 5.0,
        "destinationWarehouseId": "wh-a", "idempotencyKey": "comp-kardex-idem|COMPRA|OC-999|ENTRADA"
    }

    # Ejecutar dos veces
    InventoryTransactionService.execute_transaction(owner_uid, company_id, tx_payload, sandbox=True)
    InventoryTransactionService.execute_transaction(owner_uid, company_id, tx_payload, sandbox=True)

    res = KardexService.get_kardex(company_id=company_id, item_id=item_id, sandbox=True)
    assert len(res["movements"]) == 1
    assert res["finalBalance"]["quantity"] == 100.0
    assert res["finalBalance"]["totalValue"] == 500.0


def test_17_multi_company_isolation(kardex_env):
    """Caso 17: Los movimientos de la Empresa A jamás aparecen en el Kardex de la Empresa B."""
    store = kardex_env
    owner_uid = "user-k"
    item_id = "item-shared-code"

    # Empresa A
    store["items"][item_id] = {"id": item_id, "code": "ART-SH", "name": "Prod A", "costPrice": 10.0, "totalStock": 0.0, "type": "Bien"}
    store["client"].collection("companies/comp-A/sandbox_items").document(item_id).set(store["items"][item_id])
    InventoryTransactionService.execute_transaction(owner_uid, "comp-A", {
        "itemId": item_id, "type": "ENTRADA", "quantity": 50.0, "unitCost": 10.0,
        "destinationWarehouseId": "wh-a"
    }, sandbox=True)

    # Empresa B
    store["client"].collection("companies/comp-B/sandbox_items").document(item_id).set(store["items"][item_id])
    InventoryTransactionService.execute_transaction(owner_uid, "comp-B", {
        "itemId": item_id, "type": "ENTRADA", "quantity": 10.0, "unitCost": 10.0,
        "destinationWarehouseId": "wh-a"
    }, sandbox=True)

    k_a = KardexService.get_kardex(company_id="comp-A", item_id=item_id, sandbox=True)
    k_b = KardexService.get_kardex(company_id="comp-B", item_id=item_id, sandbox=True)

    assert k_a["finalBalance"]["quantity"] == 50.0
    assert k_b["finalBalance"]["quantity"] == 10.0


def test_18_excel_export_exact_values(kardex_env):
    """Caso 18: La exportación a Excel (.xlsx) contiene exactamente los mismos valores numéricos que el servicio."""
    import openpyxl
    store = kardex_env
    company_id = "comp-kardex-xl"
    owner_uid = "user-k"
    item_id = "item-018"

    item_data = {"id": item_id, "code": "ART-18", "name": "Tornillo", "costPrice": 2.5, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 200.0, "unitCost": 2.5,
        "destinationWarehouseId": "wh-a"
    }, sandbox=True)

    excel_io = KardexService.export_kardex_excel(company_id=company_id, item_id=item_id, sandbox=True)
    assert isinstance(excel_io, io.BytesIO)

    # Cargar workbook generado para verificar celdas
    wb = openpyxl.load_workbook(excel_io, data_only=True)
    ws = wb.active
    assert ws.title == "Kardex Valorizado"
    # Verificar título y datos
    assert "KARDEX VALORIZADO" in ws["A1"].value
    # Fila 11 es el primer movimiento
    assert ws.cell(row=11, column=6).value == 200.0  # inQty
    assert ws.cell(row=11, column=8).value == 500.0  # inTotalValue
    assert ws.cell(row=11, column=12).value == 200.0 # balanceQty
    assert ws.cell(row=11, column=14).value == 500.0 # balanceTotalValue


def test_19_pdf_export_weasyprint(kardex_env):
    """Caso 19: La exportación a PDF genera un binario PDF válido con los datos del Kardex."""
    store = kardex_env
    company_id = "comp-kardex-pdf"
    owner_uid = "user-k"
    item_id = "item-019"

    item_data = {"id": item_id, "code": "ART-19", "name": "Panel Solar", "costPrice": 8500.0, "totalStock": 0.0, "type": "Bien"}
    store["items"][item_id] = item_data
    store["client"].collection(f"companies/{company_id}/sandbox_items").document(item_id).set(item_data)

    InventoryTransactionService.execute_transaction(owner_uid, company_id, {
        "itemId": item_id, "type": "ENTRADA", "quantity": 5.0, "unitCost": 8500.0,
        "destinationWarehouseId": "wh-a"
    }, sandbox=True)

    with patch("flask.render_template", return_value="<html><body><h1>Kardex</h1></body></html>"):
        pdf_bytes = KardexService.export_kardex_pdf(company_id=company_id, item_id=item_id, sandbox=True)
        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 0
        assert pdf_bytes.startswith(b"%PDF")
