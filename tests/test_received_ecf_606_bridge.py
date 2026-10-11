import io
import pytest
from unittest.mock import patch, MagicMock
from openpyxl import load_workbook

from app.services.receptor_xml_service import ReceptorXmlService
from app.services.received_ecf_bridge_service import ReceivedEcfBridgeService
from app.web.reports_606 import _to_dgii_date, _resolve_montos, _generate_606_xlsx, _generate_606_txt

SAMPLE_ECF_E31_XML = """<?xml version="1.0" encoding="utf-8"?>
<ECF xmlns="http://www.dgii.gov.do/ecf">
  <Encabezado>
    <IdDoc>
      <TipoeCF>31</TipoeCF>
      <eNCF>E310000000001</eNCF>
      <FechaVencimientoSecuencia>31-12-2026</FechaVencimientoSecuencia>
      <TipoPago>2</TipoPago>
    </IdDoc>
    <Emisor>
      <RNCEmisor>132109122</RNCEmisor>
      <RazonSocialEmisor>PROVEEDOR TEST SRL</RazonSocialEmisor>
      <DireccionEmisor>Av. Winston Churchill 100</DireccionEmisor>
      <FechaEmision>15-10-2026</FechaEmision>
    </Emisor>
    <Comprador>
      <RNCComprador>131880681</RNCComprador>
      <RazonSocialComprador>EMPRESA COMPRADORA SRL</RazonSocialComprador>
    </Comprador>
    <Totales>
      <MontoGravadoTotal>1000.00</MontoGravadoTotal>
      <TotalITBIS>180.00</TotalITBIS>
      <MontoTotal>1180.00</MontoTotal>
    </Totales>
  </Encabezado>
  <DetallesItems>
    <Item>
      <NumeroLinea>1</NumeroLinea>
      <NombreItem>Servicio de Consultoria</NombreItem>
      <IndicadorFacturacion>1</IndicadorFacturacion>
      <IndicadorBienoServicio>2</IndicadorBienoServicio>
      <CantidadItem>1</CantidadItem>
      <PrecioUnitarioItem>600.00</PrecioUnitarioItem>
      <MontoItem>600.00</MontoItem>
    </Item>
    <Item>
      <NumeroLinea>2</NumeroLinea>
      <NombreItem>Material de Oficina</NombreItem>
      <IndicadorFacturacion>1</IndicadorFacturacion>
      <IndicadorBienoServicio>1</IndicadorBienoServicio>
      <CantidadItem>1</CantidadItem>
      <PrecioUnitarioItem>400.00</PrecioUnitarioItem>
      <MontoItem>400.00</MontoItem>
    </Item>
  </DetallesItems>
</ECF>"""

SAMPLE_ECF_E34_NC_XML = """<?xml version="1.0" encoding="utf-8"?>
<ECF xmlns="http://www.dgii.gov.do/ecf">
  <Encabezado>
    <IdDoc>
      <TipoeCF>34</TipoeCF>
      <eNCF>E340000000005</eNCF>
      <TipoPago>2</TipoPago>
    </IdDoc>
    <Emisor>
      <RNCEmisor>00112345678</RNCEmisor>
      <RazonSocialEmisor>JUAN PEREZ (PERSONA FISICA)</RazonSocialEmisor>
      <FechaEmision>2026-10-20</FechaEmision>
    </Emisor>
    <Comprador>
      <RNCComprador>131880681</RNCComprador>
      <RazonSocialComprador>EMPRESA COMPRADORA SRL</RazonSocialComprador>
    </Comprador>
    <Totales>
      <MontoGravadoTotal>500.00</MontoGravadoTotal>
      <TotalITBIS>90.00</TotalITBIS>
      <MontoTotal>590.00</MontoTotal>
    </Totales>
  </Encabezado>
  <InformacionReferencia>
    <NCFModificado>E310000000001</NCFModificado>
  </InformacionReferencia>
  <DetallesItems>
    <Item>
      <NumeroLinea>1</NumeroLinea>
      <NombreItem>Devolución Mercancía</NombreItem>
      <IndicadorFacturacion>1</IndicadorFacturacion>
      <IndicadorBienoServicio>1</IndicadorBienoServicio>
      <CantidadItem>1</CantidadItem>
      <PrecioUnitarioItem>500.00</PrecioUnitarioItem>
      <MontoItem>500.00</MontoItem>
    </Item>
  </DetallesItems>
</ECF>"""


def test_parse_ecf_detail_extracts_606_fields():
    detail, error = ReceptorXmlService.parse_ecf_detail(SAMPLE_ECF_E31_XML.encode("utf-8"))
    assert error is None
    assert detail["encf"] == "E310000000001"
    assert detail["rnc_emisor"] == "132109122"
    assert detail["fecha_emision_iso"] == "2026-10-15"
    assert detail["monto_servicios"] == 600.00
    assert detail["monto_bienes"] == 400.00
    assert detail["monto_total"] == 1180.00
    assert detail["total_itbis"] == 180.00
    assert detail["forma_pago_606"] == "02"


def test_parse_ecf_detail_e34_ncf_modificado():
    detail, error = ReceptorXmlService.parse_ecf_detail(SAMPLE_ECF_E34_NC_XML.encode("utf-8"))
    assert error is None
    assert detail["encf"] == "E340000000005"
    assert detail["rnc_emisor"] == "00112345678"
    assert detail["ncf_modificado"] == "E310000000001"
    assert detail["fecha_emision_iso"] == "2026-10-20"
    assert detail["monto_bienes"] == 500.00


def test_to_dgii_date():
    assert _to_dgii_date("2026-10-15") == "20261015"
    assert _to_dgii_date("15-10-2026") == "20261015"
    assert _to_dgii_date("2026/10/15") == "20261015"
    assert _to_dgii_date("15/10/2026") == "20261015"


def test_generate_606_xlsx_conforms_to_dgii_template():
    expenses = [
        {
            "rncEmisor": "132109122",  # 9 dígitos -> Tipo Id = 1
            "tipoGastoDGII": "02",
            "ncf": "E310000000001",
            "ncfModificado": "",
            "date": "2026-10-15",
            "dueDate": "",
            "montoServicios": 600.00,
            "montoBienes": 400.00,
            "amount": 1000.00,
            "itbisAmount": 180.00,
            "itbisLlevadoCosto": 0.00,
            "formaPago": "02",
        },
        {
            "rncEmisor": "00112345678",  # 11 dígitos (cédula) -> Tipo Id = 2
            "tipoGastoDGII": "02",
            "ncf": "E340000000005",
            "ncfModificado": "E310000000001",
            "date": "2026-10-20",
            "dueDate": "",
            "montoServicios": 0.00,
            "montoBienes": 500.00,
            "amount": 500.00,
            "itbisAmount": 90.00,
            "itbisLlevadoCosto": 0.00,
            "formaPago": "02",
        },
    ]

    dest = _generate_606_xlsx("131880681", "202610", expenses)
    dest.seek(0)
    wb = load_workbook(dest)
    ws = wb.active

    assert ws.title == "Herramienta Formato 606"
    assert ws["C4"].value == "131880681"
    assert ws["C5"].value == "202610"
    assert ws["C6"].value == 2

    # Fila 12 (registro 1: RNC juridico)
    assert ws.cell(row=12, column=2).value == "132109122"
    assert ws.cell(row=12, column=3).value == "1"  # Tipo Id = 1
    assert ws.cell(row=12, column=5).value == "E310000000001"
    assert ws.cell(row=12, column=7).value == "20261015"
    assert ws.cell(row=12, column=11).value == 600.00
    assert ws.cell(row=12, column=12).value == 400.00
    assert ws.cell(row=12, column=13).value == 1000.00
    assert ws.cell(row=12, column=14).value == 180.00

    # Fila 13 (registro 2: Persona fisica con cedula y NC)
    assert ws.cell(row=13, column=2).value == "00112345678"
    assert ws.cell(row=13, column=3).value == "2"  # Tipo Id = 2 (cedula)
    assert ws.cell(row=13, column=5).value == "E340000000005"
    assert ws.cell(row=13, column=6).value == "E310000000001"  # NCF modificado
    assert ws.cell(row=13, column=7).value == "20261020"


def test_generate_606_txt_structure():
    expenses = [
        {
            "rncEmisor": "132109122",
            "tipoGastoDGII": "02",
            "ncf": "E310000000001",
            "ncfModificado": "",
            "date": "2026-10-15",
            "dueDate": "",
            "montoServicios": 600.00,
            "montoBienes": 400.00,
            "amount": 1000.00,
            "itbisAmount": 180.00,
            "itbisLlevadoCosto": 0.00,
            "formaPago": "02",
        }
    ]
    txt_content = _generate_606_txt("131880681", "202610", expenses)
    lines = txt_content.strip().split("\n")
    assert len(lines) == 2
    assert lines[0] == "606|131880681|202610|1"

    row_fields = lines[1].split("|")
    assert len(row_fields) == 23
    assert row_fields[0] == "132109122"
    assert row_fields[1] == "1"
    assert row_fields[2] == "02"
    assert row_fields[3] == "E310000000001"
    assert row_fields[5] == "20261015"
    assert row_fields[7] == "600.00"
    assert row_fields[8] == "400.00"
    assert row_fields[9] == "1000.00"
    assert row_fields[10] == "180.00"
    assert row_fields[22] == "02"


@patch("app.services.received_ecf_bridge_service.DatabaseService.save_expense")
@patch("app.services.received_ecf_bridge_service.SupplierService.get_or_create_supplier")
@patch("app.services.received_ecf_bridge_service.ReceptorRepository._resolve_collection")
def test_bridge_process_received_ecf(mock_coll, mock_supplier, mock_save_expense):
    mock_supplier.return_value = ("supp-123", False)
    mock_doc_ref = MagicMock()
    mock_coll.return_value.document.return_value = mock_doc_ref

    ecf_doc = {
        "id": "rec-ecf-001",
        "estado_arecf": "0",
        "status": "recibido",
        "ecf_type": "31",
        "encf": "E310000000001",
        "sender_rnc": "132109122",
        "sender_name": "PROVEEDOR TEST SRL",
        "monto_total": 1180.00,
        "xml_content": SAMPLE_ECF_E31_XML,
        "received_at": "2026-10-15T10:00:00Z",
    }

    with patch.object(ReceivedEcfBridgeService, "_find_existing_expense", return_value=None):
        ok, expense_id = ReceivedEcfBridgeService.process_received_ecf("owner-999", ecf_doc, sandbox=True)

    assert ok is True
    assert expense_id is not None
    assert mock_save_expense.called
    saved_payload = mock_save_expense.call_args[0][2]
    assert saved_payload["encf"] == "E310000000001"
    assert saved_payload["rncEmisor"] == "132109122"
    assert saved_payload["includeIn606"] is True
    assert saved_payload["montoServicios"] == 600.00
    assert saved_payload["montoBienes"] == 400.00
    assert saved_payload["amount"] == 1180.00
    assert saved_payload["itbisAmount"] == 180.00


@patch("app.services.received_ecf_bridge_service.DatabaseService.save_expense")
@patch("app.services.received_ecf_bridge_service.SupplierService.get_or_create_supplier")
@patch("app.services.received_ecf_bridge_service.ReceptorRepository._resolve_collection")
def test_historical_ecf_from_previous_month_preserves_emission_date(mock_coll, mock_supplier, mock_save_expense):
    """Comprueba que un e-CF de meses anteriores conserva su fecha de emisión XML para el 606."""
    mock_supplier.return_value = ("supp-123", False)
    mock_doc_ref = MagicMock()
    mock_coll.return_value.document.return_value = mock_doc_ref

    historical_xml = SAMPLE_ECF_E31_XML.replace("<FechaEmision>15-10-2026</FechaEmision>", "<FechaEmision>20-07-2026</FechaEmision>")
    ecf_doc = {
        "id": "rec-ecf-hist-001",
        "estado_arecf": "0",
        "status": "recibido",
        "ecf_type": "31",
        "encf": "E310000000099",
        "sender_rnc": "132109122",
        "sender_name": "PROVEEDOR TEST SRL",
        "monto_total": 1180.00,
        "xml_content": historical_xml,
        "received_at": "2026-07-20T10:00:00Z",
    }

    with patch.object(ReceivedEcfBridgeService, "_find_existing_expense", return_value=None):
        ok, expense_id = ReceivedEcfBridgeService.process_received_ecf("owner-999", ecf_doc, sandbox=True)

    assert ok is True
    assert mock_save_expense.called
    saved_payload = mock_save_expense.call_args[0][2]
    # La fecha debe ser 2026-07-20 (Julio 2026), garantizando que se muestre en el 606 de Julio
    assert saved_payload["date"] == "2026-07-20"
    assert saved_payload["date"][:7] == "2026-07"

